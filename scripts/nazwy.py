"""Dawne nazwy terenowe z Państwowego Rejestru Nazw Geograficznych (PRNG, GUGiK).
Wybiera obiekty fizjograficzne, których nazwa (albo nazwa historyczna) wskazuje dawne grodziska, kurhany,
pola bitew, karczmy, młyny, huty, dwory, brody, gościńce itd. Wynik: data/nazwy.json."""
import io, json, re, zipfile, datetime, urllib.request, collections
import openpyxl

URL = "https://opendata.geoportal.gov.pl/prng/PRNG_OBIEKTY_FIZJOGRAFICZNE_XLSX.zip"

# kategoria -> wzorzec w nazwie (małe litery); kolejność ma znaczenie
CATS = [
    ("obronne", r"grodzisk|horodysz|\bgród|grodek|grodziec|zamczysk|zamkow|\bzam[eo]k|okop|szaniec|szańc|redut|\bwały?\b|wałowisk|baszt|strażnic"),
    ("pochowki", r"kurhan|mogił|mogiłk|kopiec|cmentarzysk|żalnik|cmentarz (szwedz|tatar|francus|choler|morow|wojen|stary|dawny|kozac|turec|żołnier)|cmentarzyk|trupi|umarł|nieboszcz"),
    ("wojna", r"szwedz|tatar|francuz|francusk|krzyżac|turec|kozac|moskal|rosyj|napoleon|bitw|pobojowisk|krwaw|\bobóz|obozowisk|wojsk|husar|rycer|bitewn|szubienic"),
    ("skarby", r"skarb|złot[aey]? (góra|górka|dół|łąka|pole|potok)|srebrn|pieniądz|dukat|talar|zbój|zbójeck|rozbójnik"),
    ("gospodarcze", r"karczm|krczm|austeri|zajazd|\bmłyn|młynisk|młynówk|wiatrak|folusz|tartak|\bhut[ay]?\b|huciska|hucisk|smolar|smolnik|potaż|potasz|rudnik|rudnisk|kuźni|kużnic|cegielni|wapienni|browar|gorzelni|olejarni|piec|dziegciar|węglar|mielerz"),
    ("osadnicze", r"dworzysk|\bdw[oó]r\b|dworek|dworski|folwark|kościelisk|cerkiewisk|kaplic|klasztor|pustelni|\bzb[oó]r|siedlisk|grodzisko|osadzisk|wsi?sko|staw[ay]? dworsk|plebani|zamczysko|pałac|stare siedlisko|stara wieś"),
    ("drogi", r"gościniec|gościńc|\btrakt|stara droga|dawna droga|\bbr[oó]d\b|przewóz|przeprawa|grobla|\bmost\b|carska|cesarsk|królewsk[aiey] droga|szosa|pocztow"),
    ("kamienie", r"kamień|kamienie|głaz"),
]
CATS = [(k, re.compile(v)) for k, v in CATS]


def col(head, *keys):
    for i, h in enumerate(head):
        hl = (h or "").lower()
        if all(k in hl for k in keys):
            return i
    return None


def parse_coords(v):
    """Zwraca (lat, lon) z tekstu w stopniach dziesiętnych albo w stopniach, minutach i sekundach."""
    if v is None: return None
    t = str(v).replace(",", ".")
    dms = re.findall(r"(\d+)\s*°\s*(\d+)\s*['’′]\s*([\d.]+)", t)
    if len(dms) >= 2:
        vals = [int(a) + int(b) / 60 + float(c) / 3600 for a, b, c in dms[:2]]
    else:
        nums = [float(x) for x in re.findall(r"-?\d+\.\d+", t)]
        if len(nums) < 2: return None
        vals = nums[:2]
    a, b = vals
    if 49 <= a <= 55 and 14 <= b <= 24.5: return a, b
    if 49 <= b <= 55 and 14 <= a <= 24.5: return b, a
    return None


raw = urllib.request.urlopen(urllib.request.Request(URL, headers={"User-Agent": "mapa-detektorysty/1.0"}), timeout=900).read()
z = zipfile.ZipFile(io.BytesIO(raw))
name = [n for n in z.namelist() if n.lower().endswith(".xlsx") and "legenda" not in n.lower()][0]
wb = openpyxl.load_workbook(io.BytesIO(z.read(name)), read_only=True)
ws = wb[wb.sheetnames[0]]
rows = ws.iter_rows(values_only=True)
head = [str(h or "") for h in next(rows)]
C = {
    "name": col(head, "nazwa g"), "hist": col(head, "historyczn"), "obo": col(head, "oboczn"),
    "kind": col(head, "rodzaj obiektu"), "cat": col(head, "kategoria"), "info": col(head, "informacje"),
    "msc": col(head, "miejscowo", "nadrz"), "gm": col(head, "nazwa gminy"), "pw": col(head, "powiatu"),
    "wj": col(head, "wojew"), "xy": col(head, "wgs"), "repr": col(head, "rodzaj punktu"),
}
probe = {"file": name, "header": head, "cols": C, "sample": [], "kinds": collections.Counter()}
g = lambda r, k: (str(r[C[k]]).strip() if C[k] is not None and C[k] < len(r) and r[C[k]] not in (None, "") else "")

out, seen, n = [], set(), 0
for r in rows:
    n += 1
    if n <= 5: probe["sample"].append([str(x)[:80] for x in r])
    kind = g(r, "kind"); probe["kinds"][kind] += 1
    nm, hist, obo = g(r, "name"), g(r, "hist"), g(r, "obo")
    text = " ".join([nm, hist, obo]).lower()
    cat = next((k for k, rx in CATS if rx.search(text)), None)
    if not cat: continue
    if cat == "kamienie" and "pomnik" not in kind.lower() and "głaz" not in kind.lower(): continue
    ll = parse_coords(g(r, "xy"))
    if not ll: continue
    key = (nm, round(ll[0], 3), round(ll[1], 3))
    if key in seen: continue
    seen.add(key)
    out.append({"n": nm, "h": hist, "o": obo, "r": kind, "c": cat, "lat": round(ll[0], 5), "lon": round(ll[1], 5),
                "m": g(r, "msc"), "g": g(r, "gm"), "p": g(r, "pw"), "w": g(r, "wj"), "i": g(r, "info")[:200]})

probe["rows"] = n
probe["kinds"] = probe["kinds"].most_common(80)
probe["found"] = collections.Counter(o["c"] for o in out).most_common()
json.dump(probe, open("data/_prng_probe.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
if len(out) < 200:
    raise SystemExit(f"Podejrzanie mało nazw: {len(out)} (sprawdź data/_prng_probe.json)")
json.dump({"generated": datetime.date.today().isoformat(), "count": len(out), "elements": out},
          open("data/nazwy.json", "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
print("nazwy:", len(out), probe["found"])
